# -*- coding: utf-8 -*-
"""模板填充模块：将生成内容写入 docx 工作总结模板并输出成品。"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from .generator import IMAGE_NOTE, SummaryContent, cn_number
from .template_parser import TemplateSpec, parse_template

TITLE_STYLE = "Title"
H1_STYLE = "Heading 1"
H2_STYLE = "Heading 2"
BODY_STYLE = "Body Text"

# 模板格式 → 输出格式映射（pdf/无模板 → docx）
TEMPLATE_OUT_EXT = {
    ".docx": ".docx", ".md": ".md", ".txt": ".txt", ".csv": ".txt",
    ".xlsx": ".xlsx", ".xls": ".xlsx", ".pdf": ".docx",
}


def _style_name(p) -> str:
    try:
        return p.style.name if p.style is not None else ""
    except Exception:
        return ""


def _set_text(p, text: str) -> None:
    """替换段落文字，保留段落样式。"""
    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    p.add_run(text)


def _delete_para(p) -> None:
    el = p._element
    el.getparent().remove(el)


def _get_style(doc, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return None


class _BlockWriter:
    """在「章节标题之后、边界段落之前」的块内按序写入段落。"""

    def __init__(self, doc, boundary):
        self.doc = doc
        self.boundary = boundary  # 下一个 Heading 1 段落，或 None
        self.cursor = None  # 最近写入的段落

    def write(self, text: str, style_name: str):
        style = _get_style(self.doc, style_name)
        if self.cursor is not None:
            nxt = self._next(self.cursor)
            p = nxt.insert_paragraph_before(text, style) if nxt is not None \
                else self.doc.add_paragraph(text, style)
        elif self.boundary is not None:
            p = self.boundary.insert_paragraph_before(text, style)
        else:
            p = self.doc.add_paragraph(text, style)
        self.cursor = p
        return p

    def _next(self, p):
        ps = self.doc.paragraphs
        for i, q in enumerate(ps):
            if q._p is p._p:
                return ps[i + 1] if i + 1 < len(ps) else None
        return None


def _find_heading(doc, keyword: str):
    """按关键词定位章节标题段落（Heading 1）。"""
    for p in doc.paragraphs:
        if _style_name(p).startswith("Heading 1") and keyword in p.text:
            return p
    return None


def _section_block(doc, heading_p):
    """返回 (block段落列表, 结束边界段落或None)。"""
    ps = doc.paragraphs
    start = None
    for i, q in enumerate(ps):
        if q._p is heading_p._p:
            start = i + 1
            break
    if start is None:
        return [], None
    block = []
    for q in ps[start:]:
        if _style_name(q).startswith("Heading 1"):
            return block, q
        block.append(q)
    return block, None


def _fill_body_block(doc, block, boundary, texts: list[str]) -> None:
    """用 texts 整体替换 block 段落（复用已有段落，多余删除、不足补建）。"""
    writer = _BlockWriter(doc, boundary)
    reusable = list(block)
    body_style = _get_style(doc, BODY_STYLE)
    for i, text in enumerate(texts):
        if i < len(reusable):
            _set_text(reusable[i], text)
            if body_style is not None:
                reusable[i].style = body_style
            writer.cursor = reusable[i]
        else:
            writer.write(text, BODY_STYLE)
    for p in reusable[len(texts):]:
        _delete_para(p)


def _parse_item_block(block):
    """把 block 解析为 [(h2段落, [body段落...]), ...]。"""
    items: list[tuple] = []
    current = None
    for p in block:
        if _style_name(p).startswith("Heading 2"):
            current = (p, [])
            items.append(current)
        elif current is not None:
            current[1].append(p)
    return items


def _fill_items(doc, block, boundary, sections, start_index: int,
                with_image_note: bool) -> int:
    """填充「已完成/进行中」条目块，返回下一个可用序号。"""
    placeholders = _parse_item_block(block)
    writer = _BlockWriter(doc, boundary)
    body_style = _get_style(doc, BODY_STYLE)
    idx = start_index
    for i, sec in enumerate(sections):
        heading_text = f"（{cn_number(idx)}）{sec.name}"
        body_texts = list(sec.paragraphs)
        if with_image_note:
            body_texts.append(IMAGE_NOTE)
        idx += 1
        if i < len(placeholders):
            h2, bodies = placeholders[i]
            _set_text(h2, heading_text)
            writer.cursor = h2
            for j, text in enumerate(body_texts):
                if j < len(bodies):
                    _set_text(bodies[j], text)
                    if body_style is not None:
                        bodies[j].style = body_style
                    writer.cursor = bodies[j]
                else:
                    writer.write(text, BODY_STYLE)
            for p in bodies[len(body_texts):]:
                _delete_para(p)
        else:
            writer.write(heading_text, H2_STYLE)
            for text in body_texts:
                writer.write(text, BODY_STYLE)
    # 删除多余占位条目
    for h2, bodies in placeholders[len(sections):]:
        for p in bodies:
            _delete_para(p)
        _delete_para(h2)
    return idx


def _collect_plans(content: SummaryContent) -> list[str]:
    """从进行中事项里聚合「下一步计划」句子。"""
    plans: list[str] = []
    seen: set[str] = set()
    for sec in content.wip:
        for para in sec.paragraphs:
            if not para.startswith("下一步计划："):
                continue
            text = para[len("下一步计划："):].strip()
            if not text:
                continue
            text = f"{sec.name}：{text.rstrip('。；;')}。"
            key = re.sub(r"\s+", "", text)
            if key not in seen:
                seen.add(key)
                plans.append(text)
    return plans


def _build_docx_from_spec(spec: TemplateSpec, content: SummaryContent,
                          output_path: Path) -> Path:
    """非 docx 模板：按解析出的结构从零构建 docx 成品。"""
    from docx import Document

    doc = Document()
    try:
        body_style = doc.styles[BODY_STYLE]
    except KeyError:
        body_style = None

    doc.add_paragraph(content.title or spec.title, style=TITLE_STYLE)
    has_wip_section = any(s.kind == "wip" for s in spec.sections)
    item_idx = 1
    for sec in spec.sections:
        doc.add_paragraph(sec.heading, style=H1_STYLE)
        if sec.kind == "overview":
            texts = content.overview or ["（未生成总体情况内容）"]
            for text in texts:
                doc.add_paragraph(text, style=body_style)
        elif sec.kind in ("completed", "wip"):
            if sec.kind == "completed":
                items = content.completed + ([] if has_wip_section else content.wip)
            else:
                items = content.wip
            if not items:
                doc.add_paragraph("（无）", style=body_style)
            for it in items:
                doc.add_paragraph(f"（{cn_number(item_idx)}）{it.name}", style=H2_STYLE)
                item_idx += 1
                for text in it.paragraphs:
                    doc.add_paragraph(text, style=body_style)
                if sec.kind == "completed":
                    doc.add_paragraph(IMAGE_NOTE, style=body_style)
        elif sec.kind == "plan":
            plans = _collect_plans(content)
            texts = plans or sec.hints or ["（此部分请根据实际情况补充）"]
            for text in texts:
                doc.add_paragraph(text, style=body_style)
        else:  # custom：保留模板提示文字供人工补充
            texts = sec.hints or ["（此部分请根据实际情况补充）"]
            for text in texts:
                doc.add_paragraph(text, style=body_style)
    doc.save(str(output_path))
    return output_path


def _build_docx_from_form(form, content: SummaryContent, output_path: Path) -> Path:
    """表格表单模板（如绩效 PBC）：生成填好内容的表格 docx。"""
    from docx import Document

    doc = Document()
    try:
        body_style = doc.styles[BODY_STYLE]
    except KeyError:
        body_style = None

    title = content.title or form.title
    quarter = re.search(r"\d{4}年第[一二三四]季度", content.title or "")
    if quarter:
        title = f"{form.title}（{quarter.group(0)}）"
    doc.add_paragraph(title, style=TITLE_STYLE)
    for field in form.fields:
        doc.add_paragraph(field, style=body_style)
    if form.fields:
        doc.add_paragraph("", style=body_style)

    # 只保留有实质内容的列（去掉全空合并列）
    keep = [i for i, c in enumerate(form.columns) if c]
    if not keep:
        raise ValueError("表单模板未识别出有效列")

    table = doc.add_table(rows=1, cols=len(keep))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for j, i in enumerate(keep):
        hdr[j].text = form.columns[i]

    # 权重平均分配，评价标准用模板默认值
    items = content.completed + content.wip
    n = max(len(items), 1)
    base = round(1 / n, 2)
    weights = [base] * n
    if items:
        weights[-1] = round(1 - base * (n - 1), 2)

    if not items:
        row = table.add_row().cells
        row[0].text = "（无）"
    for idx, it in enumerate(items):
        first = it.paragraphs[0] if it.paragraphs else ""
        brief = re.sub(r"\s+", " ", first).strip("。；; ")
        item_text = brief[:48] if brief.startswith(it.name) else (
            f"{it.name}：{brief[:40]}" if brief else it.name)
        values = {}
        for i, role in zip(keep, (form.col_roles[i] for i in keep)):
            if role == "index":
                values[i] = str(idx + 1)
            elif role == "item":
                values[i] = item_text
            elif role == "weight":
                values[i] = str(weights[idx])
            elif role == "std":
                values[i] = form.default_std.get(i, "")
            else:
                values[i] = ""
        row = table.add_row().cells
        for j, i in enumerate(keep):
            row[j].text = values[i]

    for line in form.footer:
        doc.add_paragraph(line, style=body_style)
    doc.save(str(output_path))
    return output_path


def _default_spec() -> "TemplateSpec":
    from .template_parser import TemplateSection

    return TemplateSpec(
        title="工作总结", source="builtin",
        sections=[
            TemplateSection(heading="一、总体情况", kind="overview"),
            TemplateSection(heading="二、已完成事项", kind="completed"),
            TemplateSection(heading="三、进行中事项", kind="wip"),
        ],
    )


# ---------- Markdown / TXT 渲染 ----------
def _build_md(spec: "TemplateSpec", content: SummaryContent,
              output_path: Path, plain: bool = False) -> Path:
    """按模板结构输出 Markdown（plain=True 时输出无标记的纯文本）。"""
    title = content.title or spec.title
    has_wip_section = any(s.kind == "wip" for s in spec.sections)
    lines = [f"# {title}" if not plain else title, ""]
    item_idx = 1
    for sec in spec.sections:
        lines.append(f"## {sec.heading}" if not plain else sec.heading)
        lines.append("")
        if sec.kind == "overview":
            texts = content.overview or ["（未生成总体情况内容）"]
            for text in texts:
                lines.append(text)
        elif sec.kind in ("completed", "wip"):
            if sec.kind == "completed":
                items = content.completed + ([] if has_wip_section else content.wip)
            else:
                items = content.wip
            if not items:
                lines.append("（无）")
            for it in items:
                head = f"（{cn_number(item_idx)}）{it.name}"
                lines.append(f"### {head}" if not plain else head)
                item_idx += 1
                for text in it.paragraphs:
                    lines.append(f"> {text}" if not plain else text)
                if sec.kind == "completed":
                    lines.append(f"> {IMAGE_NOTE}" if not plain else IMAGE_NOTE)
        elif sec.kind == "plan":
            plans = _collect_plans(content)
            texts = plans or sec.hints or ["（此部分请根据实际情况补充）"]
            for text in texts:
                lines.append(text)
        else:
            for text in sec.hints or ["（此部分请根据实际情况补充）"]:
                lines.append(text)
        lines.append("")
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


# ---------- Excel 表单/章节式渲染 ----------
def _build_xlsx_form(form, content: SummaryContent, output_path: Path) -> Path:
    """表格表单模板（PBC 等）→ xlsx 成品。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "生成结果"
    title = content.title or form.title
    quarter = re.search(r"\d{4}年第[一二三四]季度", content.title or "")
    if quarter:
        title = f"{form.title}（{quarter.group(0)}）"
    ws.append([title])
    for field in form.fields:
        ws.append([field])

    keep = [i for i, c in enumerate(form.columns) if c]
    ws.append([form.columns[i] for i in keep])

    items = content.completed + content.wip
    n = max(len(items), 1)
    base = round(1 / n, 2)
    weights = [base] * n
    if items:
        weights[-1] = round(1 - base * (n - 1), 2)

    if not items:
        ws.append(["（无）"] + [""] * (len(keep) - 1))
    for idx, it in enumerate(items):
        first = it.paragraphs[0] if it.paragraphs else ""
        brief = re.sub(r"\s+", " ", first).strip("。；; ")
        item_text = brief[:48] if brief.startswith(it.name) else (
            f"{it.name}：{brief[:40]}" if brief else it.name)
        row = []
        for i in keep:
            role = form.col_roles[i]
            if role == "index":
                row.append(str(idx + 1))
            elif role == "item":
                row.append(item_text)
            elif role == "weight":
                row.append(str(weights[idx]))
            elif role == "std":
                row.append(form.default_std.get(i, ""))
            else:
                row.append("")
        ws.append(row)

    for line in form.footer:
        ws.append([line])
    wb.save(str(output_path))
    return output_path


def _build_xlsx_sections(spec: "TemplateSpec", content: SummaryContent,
                         output_path: Path) -> Path:
    """章节式 Excel 模板 → xlsx 成品（单列行式，与模板同构）。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "生成结果"
    ws.append([content.title or spec.title])
    has_wip_section = any(s.kind == "wip" for s in spec.sections)
    item_idx = 1
    for sec in spec.sections:
        ws.append([])
        ws.append([sec.heading])
        if sec.kind == "overview":
            for text in content.overview or ["（未生成总体情况内容）"]:
                ws.append([text])
        elif sec.kind in ("completed", "wip"):
            if sec.kind == "completed":
                items = content.completed + ([] if has_wip_section else content.wip)
            else:
                items = content.wip
            for it in items or []:
                ws.append([f"（{cn_number(item_idx)}）{it.name}"])
                item_idx += 1
                for text in it.paragraphs:
                    ws.append([text])
                if sec.kind == "completed":
                    ws.append([IMAGE_NOTE])
        elif sec.kind == "plan":
            for text in _collect_plans(content) or sec.hints or ["（此部分请根据实际情况补充）"]:
                ws.append([text])
        else:
            for text in sec.hints or ["（此部分请根据实际情况补充）"]:
                ws.append([text])
    wb.save(str(output_path))
    return output_path


def fill_template(template_path: str | Path | None, content: SummaryContent,
                  output_path: str | Path) -> Path:
    """按模板生成成品，输出格式跟随模板格式；模板为空时用内置标准结构输出 docx。
    返回实际输出路径（扩展名可能与传入不同）。"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # 无模板：内置标准结构 → docx
    if template_path is None:
        return _build_docx_from_spec(_default_spec(), content,
                                     output_path.with_suffix(".docx"))

    template_path = Path(template_path)
    ext = template_path.suffix.lower()
    out_ext = TEMPLATE_OUT_EXT.get(ext, ".docx")
    output_path = output_path.with_suffix(out_ext)

    # 非 docx 模板：解析结构后按模板格式渲染
    if ext != ".docx":
        spec = parse_template(template_path)
        if ext in (".xlsx", ".xls"):
            if spec.form is not None:
                return _build_xlsx_form(spec.form, content, output_path)
            return _build_xlsx_sections(spec, content, output_path)
        if spec.form is not None:
            return _build_docx_from_form(spec.form, content,
                                         output_path.with_suffix(".docx"))
        if ext == ".md":
            return _build_md(spec, content, output_path)
        if ext in (".txt", ".csv"):
            return _build_md(spec, content, output_path, plain=True)
        # pdf 等其它模板 → docx
        return _build_docx_from_spec(spec, content,
                                     output_path.with_suffix(".docx"))

    doc = Document(str(template_path))

    # 标题
    for p in doc.paragraphs:
        if _style_name(p) == TITLE_STYLE:
            _set_text(p, content.title)
            break

    # 一、总体情况
    h1 = _find_heading(doc, "总体情况")
    if h1 is not None:
        block, boundary = _section_block(doc, h1)
        _fill_body_block(doc, block, boundary, content.overview)

    next_idx = 1
    # 二、已完成事项
    h2_done = _find_heading(doc, "已完成事项")
    if h2_done is not None:
        block, boundary = _section_block(doc, h2_done)
        next_idx = _fill_items(doc, block, boundary, content.completed,
                               next_idx, with_image_note=True)

    # 三、进行中事项
    h2_wip = _find_heading(doc, "进行中事项")
    if h2_wip is not None:
        block, boundary = _section_block(doc, h2_wip)
        if content.wip:
            _fill_items(doc, block, boundary, content.wip, next_idx,
                        with_image_note=False)
        else:
            for p in block:
                _delete_para(p)
            _insert_keep(doc, h2_wip, "本季度暂无明确进行中事项。")

    # docx 但没找到标准章节标题 → 退回结构化构建
    if h1 is None and h2_done is None and h2_wip is None:
        spec = parse_template(template_path)
        return _build_docx_from_spec(spec, content, output_path.with_suffix(".docx"))

    doc.save(str(output_path))
    return output_path


def _insert_keep(doc, heading_p, text: str):
    ps = doc.paragraphs
    for i, q in enumerate(ps):
        if q._p is heading_p._p:
            nxt = ps[i + 1] if i + 1 < len(ps) else None
            style = _get_style(doc, BODY_STYLE)
            if nxt is not None:
                nxt.insert_paragraph_before(text, style)
            else:
                doc.add_paragraph(text, style)
            return
