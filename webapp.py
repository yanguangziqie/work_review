#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""工作总结生成系统（日报 / 周报 / 月报 / 季度总结）- Web 前端

启动:
    python3 webapp.py                 # 本机访问 http://127.0.0.1:8899
    python3 webapp.py --port 9000 --host 0.0.0.0   # 局域网访问
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import tempfile
import traceback
import urllib.request
import uuid
from datetime import date, datetime
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file
from werkzeug.exceptions import HTTPException

BASE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE_DIR))

from qsum.extractors import extract_file  # noqa: E402
from qsum.generator import generate_content, render_markdown  # noqa: E402
from qsum.llm import LLMError, ai_generate_content  # noqa: E402
from qsum.miner import (WorkDigest, file_in_period, filter_items, mine_file,
                        parse_period, text_in_period)  # noqa: E402
from qsum.template_filler import fill_template  # noqa: E402

app = Flask(__name__, template_folder=str(BASE_DIR / "webui"))
app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024  # 200MB

WORK_DIR = Path(tempfile.mkdtemp(prefix="qsum_web_"))
RESULTS: dict[str, dict] = {}  # token -> {"out": Path, "md": Path|None, "name": str}

DEFAULT_TEMPLATE = BASE_DIR / "templates" / "工作总结模板.docx"
ALLOWED_EXT = {".docx", ".xlsx", ".xls", ".txt", ".md", ".csv", ".pdf"}
ALLOWED_TEMPLATE_EXT = {".docx", ".xlsx", ".xls", ".txt", ".md", ".csv", ".pdf"}
_SAFE_NAME = re.compile(r"[^\w\u4e00-\u9fa5.\-（）()]+")

REPORT_LABEL = {"day": "日报", "week": "周报", "month": "月报", "quarter": "季度总结"}

EXPORT_EXT = {"docx": ".docx", "md": ".md", "txt": ".txt", "xlsx": ".xlsx", "html": ".html"}

# 时间范围格式校验（仅拦截“周报却填季度格式”这类明显错配）；周报额外允许填周内某天日期
_PERIOD_RE = {
    "day": re.compile(r"^\d{4}-\d{1,2}-\d{1,2}$"),
    "week": re.compile(r"^(\d{4}-?W\d{1,2}|\d{4}-\d{1,2}-\d{1,2})$", re.I),
    "month": re.compile(r"^\d{4}-\d{1,2}$"),
    "quarter": re.compile(r"^\d{4}-?Q[1-4]$", re.I),
}
_PERIOD_HINT = {
    "day": "如 2026-10-01",
    "week": "如 2026W40（也可填周内某天 2026-10-01）",
    "month": "如 2026-10",
    "quarter": "如 2026Q3",
}


class ExportError(Exception):
    """导出格式不支持 / 缺少组件，message 直接面向用户。"""


class FreeformLLMError(Exception):
    """自由格式生成时的大模型调用错误。"""


def current_quarter() -> str:
    now = datetime.now()
    return f"{now.year}Q{(now.month - 1) // 3 + 1}"


def _safe_filename(name: str) -> str:
    return _SAFE_NAME.sub("_", Path(name).name) or "file"


def _canon_period(report_type: str | None, period_raw: str | None,
                  ref: date | None = None) -> str:
    """构造标准格式周期串（2026-10-01 / 2026W40 / 2026-10 / 2026Q3）。

    注意：qsum.llm 会用 digest.period 再调一次 parse_period，
    中文标签（“2026年第40周”）解析不了，所以这里必须给标准格式。
    """
    if period_raw:
        return period_raw
    ref = ref or datetime.now().date()
    rt = report_type or "quarter"
    if rt == "day":
        return ref.strftime("%Y-%m-%d")
    if rt == "week":
        iso = ref.isocalendar()
        return f"{iso.year}W{iso.week}"
    if rt == "month":
        return f"{ref.year}-{ref.month:02d}"
    return f"{ref.year}Q{(ref.month - 1) // 3 + 1}"


def _normalize_period(raw: str | None) -> str | None:
    """兼容多种日期写法：2026/9/21、2026.9.21、2026年9月21日 → 2026-09-21。"""
    if not raw:
        return raw
    s = raw.strip().translate(str.maketrans({'/': '-', '.': '-', '年': '-', '月': '-', '日': ''}))
    s = re.sub(r'-+$', '', s)
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m = re.match(r"^(\d{4})-(\d{1,2})$", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}"
    m = re.match(r"^(\d{4})-?[Ww](\d{1,2})$", s)
    if m:
        return f"{m.group(1)}W{int(m.group(2))}"
    m = re.match(r"^(\d{4})-?[Qq](\d)$", s)
    if m:
        return f"{m.group(1)}Q{m.group(2)}"
    return s


_DATE_FULL = re.compile(r"(?<!\d)(\d{4})[-/年.](\d{1,2})[-/月.](\d{1,2})(?!\d)")
_DATE_YM = re.compile(r"(?<!\d)(\d{4})[-/年.](\d{1,2})(?![-/.\d]|月\d)")
_DATE_MD = re.compile(r"(?<!\d)(\d{1,2})[-/月.](\d{1,2})(?![-/.\d]|月\d)")
_DATE_YYMMDD = re.compile(r"(?<!\d)(\d{2})(\d{2})(\d{2})(?!\d)")


def _collect_dates(texts: dict[str, str]) -> tuple[list[date], list[date]]:
    """从文本与文件名里收集日期。

    返回 (完整日期, 碎片日期)。有完整日期（年月日 / 文件名 yymmdd）时应忽略碎片日期
    （仅年月、仅月日），避免把 2026-09-22 拆出个 2026-09-01 把范围撑大。
    """
    full: list[date] = []
    partial: list[date] = []
    for name, text in texts.items():
        for y, mo, d in _DATE_YYMMDD.findall(name or ""):  # 文件名里的 yymmdd，如 事情列表260920_260925
            try:
                full.append(date(2000 + int(y), int(mo), int(d)))
            except ValueError:
                pass
        for y, mo, d in _DATE_FULL.findall(text or ""):
            try:
                full.append(date(int(y), int(mo), int(d)))
            except ValueError:
                pass
        for y, mo in _DATE_YM.findall(text or ""):
            try:
                partial.append(date(int(y), int(mo), 1))
            except ValueError:
                pass
        for mo, d in _DATE_MD.findall(text or ""):
            try:
                partial.append(date(datetime.now().year, int(mo), int(d)))
            except ValueError:
                pass
    return full, partial


def _detect_span(texts: dict[str, str]) -> tuple[date, date] | None:
    """识别到的日期范围 (最早, 最晚)；识别不到返回 None。"""
    full, partial = _collect_dates(texts)
    found = full or partial
    if not found:
        return None
    return (min(found), max(found))


def _anchor_date(texts: dict[str, str]) -> date | None:
    """推断目标日期：取出现次数最多的日期（并列取较晚的）。

    比“范围中位数”稳：跨天清单里，提到最多的那天更可能是用户想要的报告日期。
    """
    full, partial = _collect_dates(texts)
    found = full or partial
    if not found:
        return None
    counts: dict[date, int] = {}
    for d in found:
        counts[d] = counts.get(d, 0) + 1
    top = max(counts.values())
    return max(d for d, c in counts.items() if c == top)


def _span_hint(texts: dict[str, str], report_type: str | None) -> str:
    span = _detect_span(texts)
    if not span:
        return ""
    anchor = _anchor_date(texts)
    sug = _canon_period(report_type, None, ref=anchor or span[0])
    return (f"；文件里的日期约 {span[0]:%Y-%m-%d} ~ {span[1]:%Y-%m-%d}"
            f"，可在时间范围里填 {sug}（或其它你想要的日期）再生成")


# ---------------------------------------------------------------------------
# 导出格式转换
# ---------------------------------------------------------------------------

def _extract_text(path: Path) -> str:
    """把已生成的成品文件抽成纯文本，供各导出格式转换。"""
    ext = path.suffix.lower()
    if ext == ".docx":
        try:
            from docx import Document
        except ImportError:
            raise ExportError("读取 Word 内容需要 python-docx（pip install python-docx）")
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines)
    if ext in (".xlsx", ".xls"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ExportError("读取 Excel 内容需要 openpyxl（pip install openpyxl）")
        wb = load_workbook(str(path), data_only=True)
        out: list[str] = []
        for ws in wb.worksheets:
            out.append(f"## {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None and str(c).strip()]
                if cells:
                    out.append(" | ".join(cells))
        return "\n".join(out)
    if ext == ".pdf":
        try:
            import pdfplumber
            with pdfplumber.open(str(path)) as pdf:
                return "\n".join((p.extract_text() or "") for p in pdf.pages)
        except ImportError:
            pass
        try:
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
        except ImportError:
            raise ExportError("读取 PDF 内容需要 pdfplumber 或 pypdf（pip install pdfplumber）")
    return path.read_text(encoding="utf-8", errors="ignore")


def _md_plain(s: str) -> str:
    s = re.sub(r"^#+\s*", "", s)
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"__(.+?)__", r"\1", s)
    return s.replace("`", "")


def _to_md(text: str, out: Path) -> Path:
    out.write_text(text, encoding="utf-8")
    return out


def _to_txt(text: str, out: Path) -> Path:
    out.write_text("\n".join(_md_plain(l) for l in text.splitlines()), encoding="utf-8")
    return out


def _to_docx(text: str, out: Path) -> Path:
    try:
        from docx import Document
    except ImportError:
        raise ExportError("导出 Word 需要 python-docx（pip install python-docx）；可先导出 Markdown/TXT")
    doc = Document()
    for line in text.splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        if s.lstrip().startswith("#"):
            doc.add_heading(_md_plain(s), level=min(s.count("#", 0, 4) or 1, 4))
        elif s.lstrip().startswith(("- ", "* ")):
            doc.add_paragraph(_md_plain(s.lstrip()[2:]), style="List Bullet")
        elif "|" in s:
            doc.add_paragraph(_md_plain(s))
        else:
            doc.add_paragraph(_md_plain(s))
    doc.save(str(out))
    return out


def _to_xlsx(text: str, out: Path) -> Path:
    try:
        from openpyxl import Workbook
    except ImportError:
        raise ExportError("导出 Excel 需要 openpyxl（pip install openpyxl）；可先导出 Markdown/TXT")
    wb = Workbook()
    ws = wb.active
    ws.title = "内容"
    rows = 0
    for line in text.splitlines():
        if not line.strip():
            continue
        cells = [c.strip() for c in line.split("|")] if "|" in line else [line.strip()]
        ws.append(cells)
        rows += 1
    if not rows:
        raise ExportError("内容为空，无法导出 Excel")
    wb.save(str(out))
    return out


def _to_html(text: str, out: Path) -> Path:
    """导出为带基础样式的 HTML 网页（标题/列表/表格自动识别）。"""
    from html import escape
    body: list[str] = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        s = lines[i].rstrip()
        if not s.strip():
            i += 1
            continue
        if "|" in s:
            rows = []
            while i < len(lines) and "|" in lines[i]:
                rows.append([_md_plain(c.strip()) for c in lines[i].split("|")])
                i += 1
            head, rest = rows[0], rows[1:]
            t = ['<table><thead><tr>' + ''.join(f'<th>{escape(c)}</th>' for c in head) + '</tr></thead><tbody>']
            for r in rest:
                t.append('<tr>' + ''.join(f'<td>{escape(c)}</td>' for c in r) + '</tr>')
            t.append('</tbody></table>')
            body.append('\n'.join(t))
            continue
        if s.lstrip().startswith("#"):
            lvl = min(max(len(s.lstrip()) - len(s.lstrip().lstrip('#')), 1), 6)
            body.append(f'<h{lvl}>{escape(_md_plain(s))}</h{lvl}>')
        elif s.lstrip().startswith(("- ", "* ")):
            items = []
            while i < len(lines) and lines[i].lstrip().startswith(("- ", "* ")):
                items.append(f'<li>{escape(_md_plain(lines[i].lstrip()[2:]))}</li>')
                i += 1
            body.append('<ul>' + ''.join(items) + '</ul>')
            continue
        else:
            body.append(f'<p>{escape(_md_plain(s))}</p>')
        i += 1
    if not body:
        raise ExportError("内容为空，无法导出 HTML")
    html_doc = (
        '<!DOCTYPE html>\n<html lang="zh-CN">\n<head>\n<meta charset="utf-8">\n'
        '<title>工作总结</title>\n<style>\n'
        "body{font-family:'Microsoft YaHei','PingFang SC',sans-serif;max-width:860px;"
        "margin:24px auto;padding:0 16px;line-height:1.8;color:#1e293b;}\n"
        "table{border-collapse:collapse;width:100%;margin:12px 0;}\n"
        "th,td{border:1px solid #cbd5e1;padding:6px 10px;text-align:left;}\n"
        "th{background:#f1f5f9;}\nh1,h2,h3{color:#1e3a8a;}\n"
        '</style>\n</head>\n<body>\n' + '\n'.join(body) + '\n</body>\n</html>\n')
    out.write_text(html_doc, encoding="utf-8")
    return out


_EXPORTERS = {"md": _to_md, "txt": _to_txt, "docx": _to_docx, "xlsx": _to_xlsx, "html": _to_html}


# ---------------------------------------------------------------------------
# 自由格式生成（无模板 + 生成要求 → 直接按要求生成，不套默认模板）
# ---------------------------------------------------------------------------

def _ai_freeform(source_text: str, requirement: str, provider: str,
                 api_base: str | None, key_val: str | None, model: str) -> str:
    """直接调大模型按生成要求产出整篇文档（Markdown）。"""
    sys_p = ("你是工作总结撰写助手。请严格按用户的生成要求，基于提供的工作记录材料，"
             "直接输出最终文档（Markdown）。只使用材料中出现的事实与数据，不得编造；"
             "不要输出与文档无关的解释。")
    user_p = f"【生成要求】\n{requirement}\n\n【工作记录材料】\n{source_text}"
    messages = [{"role": "system", "content": sys_p},
                {"role": "user", "content": user_p}]
    try:
        if provider == "ollama":
            base = (api_base or "http://localhost:11434").rstrip("/")
            url = f"{base}/api/chat"
            body = {"model": model, "messages": messages, "stream": False}
            headers = {"Content-Type": "application/json"}
        else:
            base = (api_base or "https://api.openai.com/v1").rstrip("/")
            url = f"{base}/chat/completions"
            body = {"model": model, "messages": messages}
            headers = {"Content-Type": "application/json"}
            if key_val:
                headers["Authorization"] = f"Bearer {key_val}"
        req = urllib.request.Request(url, data=json.dumps(body).encode("utf-8"),
                                     headers=headers)
        with urllib.request.urlopen(req, timeout=600) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except Exception as exc:
        raise FreeformLLMError(f"调用大模型失败：{exc}（可检查接口地址/API Key/模型名）")
    if provider == "ollama":
        text = (payload.get("message") or {}).get("content") or ""
    else:
        try:
            text = payload["choices"][0]["message"]["content"] or ""
        except Exception:
            text = ""
    text = text.strip()
    if not text:
        raise FreeformLLMError("模型返回内容为空，请重试或更换模型")
    return text


# ---------------------------------------------------------------------------
# 路由
# ---------------------------------------------------------------------------

@app.errorhandler(Exception)
def _unhandled(exc):
    """兜底：任何未捕获异常都返回 JSON（含堆栈），避免前端拿到 HTML 报 'Unexpected token <'。"""
    if isinstance(exc, HTTPException):
        return jsonify(error=f"HTTP {exc.code}: {exc.description}"), exc.code
    tb = traceback.format_exc().strip().splitlines()[-12:]
    return jsonify(error=f"服务端内部错误：{exc}（详细堆栈见日志）", log=tb), 500


@app.get("/")
def index():
    return render_template("index.html", quarter=current_quarter())


@app.post("/api/ai/test")
def api_ai_test():
    """测试模型接口连通性，并拉取可用模型列表。"""
    data = request.get_json(silent=True) or {}
    provider = (data.get("provider") or "openai").lower()
    base_in = (data.get("api_base") or "").strip()
    key_in = (data.get("api_" + "key") or "").strip()
    if provider == "ollama":
        base = (base_in or "http://localhost:11434").rstrip("/")
        url = f"{base}/api/tags"
        headers = {}
    else:
        base = (base_in or "https://api.openai.com/v1").rstrip("/")
        url = f"{base}/models"
        headers = {"Authorization": f"Bearer {key_in}"} if key_in else {}
    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, timeout=8) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
        if provider == "ollama":
            models = [m.get("name") for m in payload.get("models", []) if m.get("name")]
        else:
            models = [m.get("id") for m in payload.get("data", []) if m.get("id")]
        return jsonify(ok=True, models=models,
                       message=f"✅ 连接成功，发现 {len(models)} 个可用模型")
    except Exception as exc:
        return jsonify(ok=False, models=[], message=f"❌ 连接失败：{exc}")


@app.post("/api/generate")
def api_generate():
    files = request.files.getlist("files")
    files = [f for f in files if f and f.filename]
    if not files:
        return jsonify(error="请至少上传一个工作内容文件"), 400

    report_type = (request.form.get("report_type") or "").strip() or None
    period_raw = (request.form.get("period")
                  or request.form.get("quarter") or "").strip() or None
    period_raw = _normalize_period(period_raw)

    # 时间范围格式与报告类型错配时给出明确提示（如“周报”填了季度格式）
    if report_type and period_raw:
        pat = _PERIOD_RE.get(report_type)
        if pat and not pat.match(period_raw):
            label = REPORT_LABEL.get(report_type, report_type)
            return jsonify(error=(
                f"时间范围「{period_raw}」与报告类型「{label}」不匹配："
                f"{label}请填 {_PERIOD_HINT.get(report_type, '')}，或清空自动取当前周期"
            )), 400

    title = (request.form.get("title") or "").strip() or None
    no_filter = request.form.get("no_filter") == "1"
    no_digest = request.form.get("no_digest") == "1"
    use_ai = request.form.get("use_ai") == "1"
    provider = (request.form.get("provider") or "openai").strip()
    api_base = (request.form.get("api_base") or "").strip() or None
    key_val = (request.form.get("api_" + "key") or "").strip() or None
    model = (request.form.get("model") or "").strip()
    extra_prompt = (request.form.get("extra_prompt") or "").strip() or None

    # 模板：显式选项“是否使用工作总结模板”
    use_template = (request.form.get("use_template") or "").strip()
    tpl_file = request.files.get("template")
    if use_template == "0":
        tpl_path = None  # 不使用模板 → 直接生成
    elif tpl_file and tpl_file.filename:
        ext = Path(tpl_file.filename).suffix.lower()
        if ext not in ALLOWED_TEMPLATE_EXT:
            return jsonify(error="模板格式不支持（可上传 docx/xlsx/xls/txt/md/csv/pdf）"), 400
        tpl_path = WORK_DIR / f"template_{uuid.uuid4().hex[:8]}{ext}"
        tpl_file.save(tpl_path)
    elif use_template == "1":
        return jsonify(error="已选择“使用模板”，但未上传模板文件：请上传模板，"
                             "或将模板选项改为“不使用模板”"), 400
    else:
        tpl_path = None  # 未传选项时兼容旧行为：没模板文件就不走模板

    # 无模板 + 给了生成要求 → 自由格式直接生成（不套默认模板）
    free_form = tpl_path is None and bool(extra_prompt)

    run_dir = WORK_DIR / uuid.uuid4().hex[:12]
    run_dir.mkdir(parents=True, exist_ok=True)

    # 先落盘，之后多轮解析共用
    pre_log: list[str] = []
    saved: list[tuple[str, Path]] = []
    for f in files:
        ext = Path(f.filename).suffix.lower()
        if ext not in ALLOWED_EXT:
            pre_log.append(f"[跳过] {f.filename}: 不支持的格式 {ext}")
            continue
        save_path = run_dir / _safe_filename(f.filename)
        if save_path.exists():  # 同名文件去重
            save_path = run_dir / f"{save_path.stem}_{uuid.uuid4().hex[:4]}{save_path.suffix}"
        f.save(save_path)
        saved.append((f.filename, save_path))
    if not saved:
        return jsonify(error="没有可解析的文件（格式需为 docx/xlsx/txt/md/csv/pdf）",
                       log=pre_log), 400

    def run_pass(canon: str, apply_filter: bool):
        """跑一轮 解析/过滤/挖掘。返回 (period对象|None, digest|None, log, texts)。"""
        try:
            pobj = parse_period(report_type, canon)
        except ValueError as exc:
            return None, None, [f"[错误] {exc}"], {}
        log = list(pre_log)
        digest = WorkDigest(period=canon)
        # qsum.llm 会用 digest.period 再解析一次，中文标签解析不了，强制用标准格式
        try:
            digest.period = canon
        except Exception:
            pass
        texts: dict[str, str] = {}
        for name, path in saved:
            try:
                text = extract_file(path)
            except Exception as exc:
                digest.skipped.append(f"{name}（{exc}）")
                log.append(f"[跳过] {name}: {exc}")
                continue
            texts[name] = text
            if apply_filter and not file_in_period(path, text, pobj):
                digest.skipped.append(f"{name}（不属于 {pobj.label}）")
                log.append(f"[过滤] {name}: 不属于 {pobj.label}，已跳过")
                continue
            try:
                items, metrics = mine_file(path, text)
            except Exception as exc:  # 单个文件挖事项失败不影响其它文件
                digest.skipped.append(f"{name}（解析事项失败：{exc}）")
                log.append(f"[跳过] {name}: 解析工作事项失败：{exc}")
                continue
            n_all = len(items)
            if apply_filter:
                items = filter_items(items, pobj)
                metrics = [m for m in metrics if text_in_period(m, pobj) is not False]
                if len(items) < n_all:
                    log.append(f"[过滤] {name}: {n_all - len(items)} 条事项不在 {pobj.label}，已丢弃")
            if not items and not metrics:
                digest.skipped.append(f"{name}（未识别到工作事项）")
                log.append(f"[提示] {name}: 未识别到工作事项")
                continue
            digest.files.append(name)
            digest.items_completed.extend(i for i in items if i.status == "completed")
            digest.items_wip.extend(i for i in items if i.status == "wip")
            digest.metric_lines.extend(metrics)
            n_done = sum(1 for i in items if i.status == "completed")
            n_wip = sum(1 for i in items if i.status == "wip")
            log.append(f"[解析] {name}: 已完成 {n_done} 条 / 进行中 {n_wip} 条")
        return pobj, digest, log, texts

    # 第 1 轮：按目标周期过滤（勾了“不做时间过滤”则直接全量）
    canon = _canon_period(report_type, period_raw)
    apply_filter = not no_filter
    period, digest, log, texts = run_pass(canon, apply_filter)

    # 自动降级：目标周期内没识别到事项 → 按文件日期推断周期（用户没指定时间时），
    # 并忽略时间过滤全量纳入（而不是直接报错）
    if digest is not None and not digest.all_items and apply_filter:
        span = _detect_span(texts)
        anchor = _anchor_date(texts)
        canon2, note = canon, ""
        if anchor and not period_raw:
            cand = _canon_period(report_type, None, ref=anchor)
            if cand != canon:
                canon2 = cand
                note = (f"，检测到文件日期约 {span[0]:%Y-%m-%d} ~ {span[1]:%Y-%m-%d}"
                        f"（出现最多的日期 {anchor:%Y-%m-%d}），已自动改按 {canon2} 生成")
        log.append(f"[自动] 目标周期内没有可识别的事项{note}；"
                   f"已忽略时间过滤，纳入全部文件与事项（如需严格过滤请调整时间范围后重新生成）")
        p2, d2, l2, texts = run_pass(canon2, False)
        if d2 is not None and d2.all_items:
            period, digest = p2, d2
            log += l2

    if digest is None or not digest.all_items:
        return jsonify(
            error="未从上传文件中识别到工作事项，请检查内容" + _span_hint(texts, report_type),
            log=log,
        ), 400

    if use_ai and not model and not free_form:
        return jsonify(error="AI 生成需要填写模型名称（如 gpt-4o-mini / qwen2.5:14b）"), 400
    if free_form and use_ai and not model:
        return jsonify(error="按生成要求自由生成需要填写模型名称（如 gpt-4o-mini / qwen2.5:14b）"), 400

    # 文件名不带日期：优先用用户填的标题，否则用报告类型名（日报/周报/月报/季度总结）
    base_name = _safe_filename(title) if title else REPORT_LABEL.get(report_type or "quarter", "工作总结")

    # 先离线抽取材料（也用作条目统计与摘要）
    try:
        content = generate_content(digest, title=title)
    except Exception as exc:
        log.append(f"[错误] 内容生成内部错误：{exc!r}")
        log.extend(traceback.format_exc().strip().splitlines()[-12:])
        return jsonify(
            error=f"内容生成出现内部错误：{exc}（详细堆栈见下方日志，请截图反馈）",
            log=log,
        ), 500

    if free_form:
        log.append("[自由格式] 未上传模板且填写了生成要求 → 直接按要求生成，不套默认模板")
        try:
            source_text = render_markdown(content, digest)
        except Exception as exc:
            source_text = ""
            log.append(f"[提示] 材料摘要渲染失败：{exc}")
        if use_ai:
            try:
                doc_text = _ai_freeform(source_text, extra_prompt or "",
                                        provider, api_base, key_val, model)
            except FreeformLLMError as exc:
                return jsonify(error=f"AI 生成失败：{exc}", log=log), 400
            except Exception as exc:
                log.append(f"[错误] AI 自由生成内部错误：{exc!r}")
                log.extend(traceback.format_exc().strip().splitlines()[-12:])
                return jsonify(error=f"AI 自由生成出现内部错误：{exc}（详细堆栈见日志）",
                               log=log), 500
            log.append("[AI] 已按生成要求产出自由格式文档")
        else:
            doc_text = (f"# {title or period.label}\n\n" if title else "") + source_text
            log.append("[提示] 规则模式无法理解生成要求的文字含义，已输出自由格式总结；"
                       "如需严格按要求润色请改用 AI 生成模式")
        out_path = run_dir / f"{base_name}.md"
        try:
            out_path.write_text(doc_text, encoding="utf-8")
        except Exception as exc:
            return jsonify(error=f"写入结果文件失败：{exc}", log=log), 500
    else:
        log.append("[AI] 使用大模型生成内容" if use_ai else "[规则] 使用内置规则生成内容")
        try:
            if use_ai:
                content = ai_generate_content(
                    digest, provider=provider, api_base=api_base,
                    **{"api_key": key_val},
                    model=model, title=title, extra_prompt=extra_prompt,
                )
            else:
                if extra_prompt:
                    log.append("[提示] 已上传模板，按模板结构生成；自定义要求在 AI 模式下才会被用于措辞")
        except LLMError as exc:
            return jsonify(error=f"AI 生成失败：{exc}", log=log), 400
        except Exception as exc:
            log.append(f"[错误] 内容生成内部错误：{exc!r}")
            log.extend(traceback.format_exc().strip().splitlines()[-12:])
            return jsonify(
                error=f"内容生成出现内部错误：{exc}（详细堆栈见下方日志，请截图反馈）",
                log=log,
            ), 500

        out_path = run_dir / f"{base_name}{period.suffix}"
        try:
            out_path = fill_template(tpl_path, content, out_path)
        except Exception as exc:
            log.append(f"[错误] 模板解析/填充失败：{exc!r}")
            log.extend(traceback.format_exc().strip().splitlines()[-8:])
            return jsonify(error=f"模板解析/填充失败：{exc}", log=log), 400

    out_name = out_path.name

    digest_path = None
    if not no_digest:
        try:
            digest_path = run_dir / f"{out_path.stem}_工作内容摘要.md"
            digest_path.write_text(render_markdown(content, digest), encoding="utf-8")
        except Exception as exc:
            log.append(f"[提示] 摘要生成失败（不影响主文件）：{exc}")
            digest_path = None

    token = uuid.uuid4().hex[:16]
    RESULTS[token] = {"out": out_path, "md": digest_path, "name": out_name}
    log.append(f"[生成] {out_name}：已完成 {len(content.completed)} 个条目 / "
               f"进行中 {len(content.wip)} 个条目")

    payload = {
        "token": token,
        "output_name": out_name,
        "completed": len(content.completed),
        "wip": len(content.wip),
        "has_digest": digest_path is not None,
        "free_form": free_form,
        "log": log,
    }
    return jsonify(payload)


@app.get("/api/download/<token>/<kind>")
def api_download(token: str, kind: str):
    entry = RESULTS.get(token)
    if not entry:
        return jsonify(error="结果不存在或已过期"), 404
    if kind == "md":
        path = entry["md"]
        name = Path(entry["name"]).stem + "_工作内容摘要.md"
    else:
        path = entry["out"]
        name = entry["name"]
    if path is None or not path.exists():
        return jsonify(error="文件不存在"), 404
    return send_file(path, as_attachment=True, download_name=name)


@app.get("/api/export/<token>/<fmt>")
def api_export(token: str, fmt: str):
    """按所选格式导出；不支持的格式/缺组件返回 JSON 提示（前端展示）。"""
    entry = RESULTS.get(token)
    if not entry:
        return jsonify(error="结果不存在或已过期"), 404
    fmt = (fmt or "").lower()
    src: Path = entry["out"]
    if fmt in ("orig", ""):
        return send_file(src, as_attachment=True, download_name=entry["name"])
    if fmt not in EXPORT_EXT:
        return jsonify(error=f"不支持的导出格式：{fmt}（可选：word(.docx) / md / txt / excel(.xlsx) / html）"), 400
    try:
        text = _extract_text(src)
    except ExportError as exc:
        return jsonify(error=str(exc)), 400
    except Exception as exc:
        return jsonify(error=f"读取生成内容失败：{exc}"), 500
    out = src.with_name(src.stem + f"_导出{EXPORT_EXT[fmt]}")
    try:
        out = _EXPORTERS[fmt](text, out)
    except ExportError as exc:
        return jsonify(error=str(exc)), 400
    except Exception as exc:
        return jsonify(error=f"导出 {fmt} 失败：{exc}"), 500
    name = Path(entry["name"]).stem + EXPORT_EXT[fmt]
    return send_file(out, as_attachment=True, download_name=name)


def main() -> int:
    ap = argparse.ArgumentParser(description="工作总结生成系统（日报/周报/月报/季度总结）Web 前端")
    ap.add_argument("--host", default="127.0.0.1", help="监听地址（0.0.0.0 允许局域网访问）")
    ap.add_argument("--port", type=int, default=8899, help="端口号")
    ap.add_argument("--debug", action="store_true")
    args = ap.parse_args()
    print(f"工作总结生成系统已启动（日报/周报/月报/季度总结）: http://{args.host}:{args.port}")
    app.run(host=args.host, port=args.port, debug=args.debug)
    return 0


if __name__ == "__main__":
    sys.exit(main())
